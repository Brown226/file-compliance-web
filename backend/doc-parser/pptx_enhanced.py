"""
增强型 PPTX 解析器
- 使用 python-pptx 原生解析，逐幻灯片提取文本、表格、图片
- 支持分组形状、备注、图表数据
- 生成 Markdown 格式 + 结构化数据
"""

import io
import logging
from typing import Optional

logger = logging.getLogger("markitdown-service")


def _escape_pipe(text: str) -> str:
    """转义 Markdown 表格中的管道符"""
    return text.replace("|", "&#124;")


def _clean_cell(value) -> str:
    """清理单元格内容"""
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\n", "<br>")
    text = _escape_pipe(text)
    return text


def _table_to_markdown(table) -> str:
    """python-pptx Table -> Markdown 表格"""
    rows_data = []
    for row in table.rows:
        cells = [_clean_cell(cell.text) for cell in row.cells]
        rows_data.append(cells)

    if not rows_data:
        return ""

    num_cols = max(len(r) for r in rows_data)
    for r in rows_data:
        while len(r) < num_cols:
            r.append("")

    header = rows_data[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * num_cols) + " |",
    ]
    for row in rows_data[1:]:
        lines.append("| " + " | ".join(row[:num_cols]) + " |")

    return "\n".join(lines)


def _extract_table_rows(table) -> list:
    """提取表格数据为二维数组"""
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(cells)
    return rows


def _is_picture(shape) -> bool:
    """检测形状是否为图片"""
    try:
        import pptx
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            return True
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PLACEHOLDER:
            if hasattr(shape, "image"):
                return True
    except Exception:
        pass
    return False


def _is_table(shape) -> bool:
    """检测形状是否为表格"""
    try:
        import pptx
        return shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.TABLE
    except Exception:
        return False


def _is_group(shape) -> bool:
    """检测形状是否为分组"""
    try:
        import pptx
        return shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP
    except Exception:
        return False


def _extract_shape_text(shape) -> str:
    """从形状中提取纯文本（含文本框）"""
    if shape.has_text_frame:
        return shape.text_frame.text.strip()
    return ""


def _process_shape(shape, markdown_parts: list, paragraphs: list, tables: list,
                    slide_idx: int, is_title_shape) -> None:
    """递归处理单个形状"""
    import pptx

    # 图片
    if _is_picture(shape):
        alt_text = ""
        try:
            alt_text = shape._element._nvXxPr.cNvPr.attrib.get("descr", "")
        except Exception:
            pass
        if not alt_text:
            alt_text = getattr(shape, "name", "image")
        markdown_parts.append(f"![{alt_text}]")
        return

    # 表格
    if _is_table(shape):
        table_md = _table_to_markdown(shape.table)
        if table_md:
            markdown_parts.append(table_md)
            rows = _extract_table_rows(shape.table)
            if rows:
                tables.append({"page": slide_idx, "rows": rows, "caption": None})
        return

    # 图表
    if hasattr(shape, "has_chart") and shape.has_chart:
        chart_md = _convert_chart(shape.chart)
        if chart_md:
            markdown_parts.append(chart_md)
        return

    # 文本框
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if not text:
            return

        if shape == is_title_shape:
            markdown_parts.append(f"# {text}")
            paragraphs.append({"text": text, "style": "Heading1", "page": slide_idx})
        else:
            markdown_parts.append(text)
            paragraphs.append({"text": text, "style": "Normal", "page": slide_idx})
        return

    # 分组形状：递归处理子形状
    if _is_group(shape):
        try:
            sorted_shapes = sorted(
                shape.shapes,
                key=lambda x: (
                    float("-inf") if not x.top else x.top,
                    float("-inf") if not x.left else x.left,
                ),
            )
            for sub_shape in sorted_shapes:
                _process_shape(sub_shape, markdown_parts, paragraphs, tables,
                               slide_idx, None)
        except Exception:
            pass


def _convert_chart(chart) -> str:
    """将图表转换为 Markdown 表格"""
    try:
        md = "\n\n### Chart"
        if chart.has_title:
            md += f": {chart.chart_title.text_frame.text}"
        md += "\n\n"

        data = []
        category_names = [c.label for c in chart.plots[0].categories]
        series_names = [s.name for s in chart.series]
        data.append(["Category"] + series_names)

        for idx, category in enumerate(category_names):
            row = [category]
            for series in chart.series:
                row.append(str(series.values[idx]))
            data.append(row)

        lines = []
        for i, row in enumerate(data):
            lines.append("| " + " | ".join(row) + " |")
            if i == 0:
                lines.append("|" + "|".join(["---"] * len(row)) + "|")

        return md + "\n".join(lines)
    except Exception:
        return ""


def enhanced_pptx_parse(content: bytes, filename: str) -> Optional[dict]:
    """
    增强型 PPTX 解析：使用 python-pptx 原生解析。

    Returns:
        dict: 兼容 ParseResult 格式的结果，失败返回 None。
    """
    try:
        from pptx import Presentation
    except ImportError:
        logger.warning("python-pptx 未安装，无法使用增强 PPTX 解析")
        return None

    try:
        prs = Presentation(io.BytesIO(content))
    except Exception as e:
        logger.warning(f"python-pptx 打开文件失败: {filename} - {e}")
        return None

    try:
        markdown_parts = []
        paragraphs = []
        tables = []
        has_images = False

        for slide_idx, slide in enumerate(prs.slides):
            markdown_parts.append(f"\n<!-- Slide {slide_idx + 1} -->\n")

            title_shape = slide.shapes.title

            # 按位置排序形状（从上到下、从左到右）
            try:
                sorted_shapes = sorted(
                    slide.shapes,
                    key=lambda x: (
                        float("-inf") if not x.top else x.top,
                        float("-inf") if not x.left else x.left,
                    ),
                )
            except Exception:
                sorted_shapes = list(slide.shapes)

            for shape in sorted_shapes:
                if _is_picture(shape):
                    has_images = True
                _process_shape(shape, markdown_parts, paragraphs, tables,
                               slide_idx, title_shape)

            # 备注
            if slide.has_notes_slide:
                try:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
                    if notes_text:
                        markdown_parts.append(f"\n### Notes:\n{notes_text}")
                        paragraphs.append({"text": notes_text, "style": "Notes", "page": slide_idx})
                except Exception:
                    pass

        markdown_text = "\n\n".join(markdown_parts).strip()
        plain_text = "\n\n".join(p["text"] for p in paragraphs if p["text"])

        return {
            "text": plain_text,
            "pages": [plain_text] if plain_text else [],
            "metadata": {
                "page_count": len(prs.slides),
                "has_tables": len(tables) > 0,
                "has_images": has_images,
                "parse_error": None,
            },
            "structure": {
                "paragraphs": paragraphs,
                "tables": tables,
                "headers": [],
                "dimensions": [],
            },
            "markdown": markdown_text,
        }

    except Exception as e:
        logger.error(f"增强 PPTX 解析失败: {filename} - {e}", exc_info=True)
        return None
